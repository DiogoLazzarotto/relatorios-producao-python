"""Valida produção CSV/XLSX e gera relatório HTML/JSON (XLSX requer openpyxl)."""
import argparse
import csv
import html
import json
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

CATEGORIES = ('Medicamentos', 'Premix', 'Mineral', 'Ração', 'Recebimentos')
FIELDS = ('data', 'categoria', 'produto', 'batidas', 'kg')

def number(value, field, line):
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        raise ValueError(f'Linha {line}: {field} inválido')
    try:
        n = Decimal(str(value).replace(',', '.'))
    except InvalidOperation:
        raise ValueError(f'Linha {line}: {field} inválido') from None
    if not n.is_finite() or n < 0:
        raise ValueError(f'Linha {line}: {field} deve ser finito e não negativo')
    return n


def _validate_row(values, line):
    if len(values) != len(FIELDS):
        raise ValueError(f'Linha {line}: número de colunas inválido')
    r = dict(zip(FIELDS, values))
    for field, value in r.items():
        if value is None or (isinstance(value, str) and not value.strip()):
            raise ValueError(f'Linha {line}: {field} vazio')
        if isinstance(value, str):
            r[field] = value.strip()
    value = r['data']
    try:
        if isinstance(value, datetime):
            if value.time() != time() or value.tzinfo is not None:
                raise ValueError
            day = value.date()
        elif isinstance(value, date):
            day = value
        elif isinstance(value, str):
            day = date.fromisoformat(value)
        else:
            raise ValueError
    except ValueError:
        raise ValueError(f'Linha {line}: data inválida') from None
    if (r['categoria'] not in CATEGORIES or not isinstance(r['produto'], str)
            or not r['produto']):
        raise ValueError(f'Linha {line}: categoria ou produto inválido')
    b = number(r['batidas'], 'batidas', line)
    kg = number(r['kg'], 'kg', line)
    if b != b.to_integral_value():
        raise ValueError(f'Linha {line}: batidas devem ser inteiras')
    if r['categoria'] in CATEGORIES[:3] and kg != 0:
        raise ValueError(f'Linha {line}: categorias auxiliares usam apenas batidas (kg=0)')
    if r['categoria'] == 'Recebimentos' and b != 0:
        raise ValueError(f'Linha {line}: recebimentos usam apenas kg (batidas=0)')
    return dict(data=day, categoria=r['categoria'], produto=r['produto'], batidas=int(b), kg=kg)


def _check_header(values):
    if tuple(values) != FIELDS:
        raise ValueError('Cabeçalho esperado: ' + ';'.join(FIELDS))


def _load_csv(path):
    rows = []
    with open(path, encoding='utf-8-sig', newline='') as f:
        reader = csv.DictReader(f, delimiter=';')
        _check_header(reader.fieldnames or ())
        for r in reader:
            line = reader.line_num
            if None in r:
                raise ValueError(f'Linha {line}: número de colunas inválido')
            rows.append(_validate_row(list(r.values()), line))
    return rows


def _load_xlsx(path):
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
        from openpyxl.xml import LXML
    except ImportError:
        raise ValueError('XLSX requer openpyxl. Instale com: '
                         'python -m pip install -r requirements-xlsx.txt') from None
    xml_errors = (ParseError,)
    if LXML:
        from lxml.etree import XMLSyntaxError
        xml_errors += (XMLSyntaxError,)
    try:
        # O arquivo também fecha caso o construtor do workbook falhe.
        with open(path, 'rb') as source:
            workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
            try:
                if not workbook.worksheets:
                    _check_header(())
                sheet = workbook.worksheets[0]
                # Não confiar em dimensões incorretas que ocultem linhas/colunas.
                sheet.reset_dimensions()
                iterator = sheet.iter_rows()
                header = [cell.value for cell in next(iterator, ())]
                while header and header[-1] is None:
                    header.pop()
                _check_header(header)
                rows = []
                for line, cells in enumerate(iterator, 2):
                    values = [cell.value for cell in cells]
                    if all(value is None for value in values):
                        continue
                    if any(value is not None for value in values[len(FIELDS):]):
                        raise ValueError(f'Linha {line}: número de colunas inválido')
                    for field, cell in zip(FIELDS, cells):
                        if cell.data_type in ('f', 'e'):
                            raise ValueError(f'Linha {line}: {field} contém fórmula ou erro Excel')
                    values = values[:len(FIELDS)]
                    values += [None] * (len(FIELDS) - len(values))
                    rows.append(_validate_row(values, line))
                return rows
            finally:
                workbook.close()
    except (BadZipFile, InvalidFileException, KeyError, EOFError, *xml_errors) as error:
        raise ValueError(f'Arquivo XLSX inválido: {error}') from None


def load(path):
    path = Path(path)
    if path.suffix.lower() == '.csv':
        return _load_csv(path)
    if path.suffix.lower() == '.xlsx':
        return _load_xlsx(path)
    raise ValueError('Formato não suportado: use CSV ou XLSX')

def summarize(rows, start=None, end=None):
    if start and end and start > end:
        raise ValueError('Data inicial deve ser anterior ou igual à final')
    groups = {}
    for r in rows:
        if (start and r['data'] < start) or (end and r['data'] > end):
            continue
        key = (r['categoria'], r['produto'])
        target = groups.setdefault(key, {'batidas': 0, 'kg': Decimal(0)})
        target['batidas'] += r['batidas']; target['kg'] += r['kg']
    return groups

def export(groups, output, title):
    output = Path(output); output.mkdir(parents=True, exist_ok=True)
    payload = [{'categoria': c, 'produto': p, 'batidas': v['batidas'], 'kg': str(v['kg'])}
               for (c, p), v in sorted(groups.items())]
    (output/'totais.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    sections = []
    for cat in CATEGORIES:
        items = [(p, v) for (c, p), v in sorted(groups.items()) if c == cat]
        if not items: continue
        weight = cat in ('Ração', 'Recebimentos')
        heads = '<th>Batidas</th><th>kg</th>' if cat == 'Ração' else ('<th>kg</th>' if weight else '<th>Batidas</th>')
        trs = []
        for product, v in items:
            cells = f"<td>{v['batidas']}</td><td>{v['kg']}</td>" if cat == 'Ração' else f"<td>{v['kg'] if weight else v['batidas']}</td>"
            trs.append('<tr><td>'+html.escape(product)+'</td>'+cells+'</tr>')
        total = sum((v['kg'] for _, v in items), Decimal(0)) if weight else sum(v['batidas'] for _, v in items)
        footer = f'Total: {total} kg' if weight else f'Total: {total} batidas'
        if cat == 'Ração': footer += f' · {total/Decimal(1000)} toneladas'
        sections.append(f'<section><h2>{cat}</h2><table><thead><tr><th>Produto</th>{heads}</tr></thead><tbody>'+''.join(trs)+f'</tbody></table><p><strong>{footer}</strong></p></section>')
    body = ''.join(sections) or '<p>Nenhum registro no período.</p>'
    page = '<!doctype html><html lang="pt-BR"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+html.escape(title)+'</title><style>body{font:16px system-ui;color:#172b42;max-width:900px;margin:40px auto;padding:0 24px}h1,h2{color:#124b80}section{break-inside:avoid;margin:30px 0}table{width:100%;border-collapse:collapse}th,td{padding:10px;border-bottom:1px solid #ccd8e5;text-align:left}th{background:#eaf2fa}@media print{body{margin:0;font-size:12px}}</style><h1>'+html.escape(title)+'</h1><p>Dados fictícios para demonstração · pesos informados, sem conversão de batidas.</p>'+body+'</html>'
    (output/'relatorio.html').write_text(page, encoding='utf-8')

def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('input', help='arquivo CSV ou XLSX'); ap.add_argument('--start', type=date.fromisoformat)
    ap.add_argument('--end', type=date.fromisoformat); ap.add_argument('--output', default='output')
    args = ap.parse_args()
    try:
        groups = summarize(load(args.input), args.start, args.end)
        export(groups, args.output, f'Produção e recebimentos · {args.start or "início"} a {args.end or "fim"}')
    except (ValueError, OSError) as e:
        ap.exit(2, f'Erro: {e}\n')
    print(f'Relatório criado em {args.output}/relatorio.html')
if __name__ == '__main__': main()

