"""Contrato proposto para issue #2. Execute separadamente; XLSX ainda não implementado."""
import csv
import json
import subprocess
import sys
import tempfile
import unittest
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src'))
from report import FIELDS, load, summarize, export


class XlsxAcceptanceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)

    def workbook(self, rows, header=FIELDS, filename='input.xlsx'):
        path = self.folder / filename
        wb = Workbook()
        ws = wb.active
        ws.title = 'Produção'
        if header is not None:
            ws.append(list(header))
        for row in rows:
            ws.append(row)
        wb.save(path)
        wb.close()
        return path

    def csv_file(self, rows):
        path = self.folder / 'input.csv'
        with path.open('w', encoding='utf-8', newline='') as output:
            writer = csv.writer(output, delimiter=';')
            writer.writerow(FIELDS)
            writer.writerows(rows)
        return path

    def valid_row(self):
        return ['2026-09-21', 'Ração', 'Produto', 1, '0.1']

    def test_example_parity_and_known_totals(self):
        source = ROOT / 'examples' / 'producao.csv'
        with source.open(encoding='utf-8', newline='') as f:
            reader = csv.reader(f, delimiter=';')
            next(reader)
            raw = [row for row in reader if row]
        xlsx = self.workbook([[date.fromisoformat(d), c, p, int(b), int(k)]
                              for d, c, p, b, k in raw])
        csv_rows, xlsx_rows = load(source), load(xlsx)
        self.assertEqual(csv_rows, xlsx_rows)
        groups = summarize(xlsx_rows)
        self.assertEqual(groups, summarize(csv_rows))
        for category, batches, kg in [('Ração', 58, '145000'), ('Medicamentos', 100, '0'),
                                      ('Premix', 100, '0'), ('Mineral', 18, '0'),
                                      ('Recebimentos', 0, '36720')]:
            values = [v for (c, _), v in groups.items() if c == category]
            self.assertEqual(sum(v['batidas'] for v in values), batches)
            self.assertEqual(sum((v['kg'] for v in values), Decimal(0)), Decimal(kg))
        self.assertEqual(summarize(xlsx_rows, date(2026, 9, 22), date(2026, 9, 22)),
                         summarize(csv_rows, date(2026, 9, 22), date(2026, 9, 22)))

    def test_decimal_and_exports_match_csv(self):
        raw = [['2026-09-21', 'Ração', '<script>alert(1)</script>', 1, value]
               for value in ['0.1', '0.2', '0,3', '123456789012345.123456789']]
        csv_groups = summarize(load(self.csv_file(raw)))
        xlsx_groups = summarize(load(self.workbook(raw)))
        self.assertEqual(xlsx_groups, csv_groups)
        total = xlsx_groups[('Ração', '<script>alert(1)</script>')]['kg']
        self.assertIsInstance(total, Decimal)
        self.assertEqual(total, Decimal('123456789012345.723456789'))
        for label, groups in [('csv', csv_groups), ('xlsx', xlsx_groups)]:
            export(groups, self.folder / label, '<Título>')
        for filename in ['relatorio.html', 'totais.json']:
            self.assertEqual((self.folder / 'csv' / filename).read_bytes(),
                             (self.folder / 'xlsx' / filename).read_bytes())
        payload = json.loads((self.folder / 'xlsx' / 'totais.json').read_text(encoding='utf-8'))
        self.assertEqual(payload[0]['kg'], str(total))
        html = (self.folder / 'xlsx' / 'relatorio.html').read_text(encoding='utf-8')
        self.assertIn('&lt;script&gt;', html)
        self.assertNotIn('<script>', html)
        self.assertIn('&lt;Título&gt;', html)
        self.assertIn(f'{total / Decimal(1000)} toneladas', html)

    def test_numeric_cells_use_decimal_without_float_artifacts(self):
        rows = [self.valid_row(), self.valid_row()]
        rows[0][4], rows[1][4] = 0.1, 0.2
        loaded = load(self.workbook(rows))
        self.assertTrue(all(isinstance(r['kg'], Decimal) for r in loaded))
        self.assertEqual(summarize(loaded)[('Ração', 'Produto')]['kg'], Decimal('0.3'))

    def test_date_cells_and_iso_text(self):
        for value in ['2026-09-21', date(2026, 9, 21), datetime(2026, 9, 21)]:
            with self.subTest(value=value):
                row = self.valid_row()
                row[0] = value
                self.assertEqual(load(self.workbook([row]))[0]['data'], date(2026, 9, 21))

    def test_invalid_dates(self):
        for value in ['2026-02-30', '21/09/2026', 46286, datetime(2026, 9, 21, 12)]:
            with self.subTest(value=value):
                row = self.valid_row()
                row[0] = value
                with self.assertRaisesRegex(ValueError, r'Linha 2:.*data'):
                    load(self.workbook([row]))

    def test_headers_are_exact(self):
        headers = [None, FIELDS[:-1], (*FIELDS, 'extra'),
                   ('categoria', 'data', 'produto', 'batidas', 'kg'),
                   ('data', 'categoria', 'produto', 'batidas', 'batidas'),
                   ('Data', *FIELDS[1:]), (' data', *FIELDS[1:])]
        for header in headers:
            with self.subTest(header=header):
                with self.assertRaisesRegex(ValueError, 'Cabeçalho'):
                    load(self.workbook([], header=header))

    def test_empty_required_cells_and_whitespace(self):
        for index in range(5):
            for value in [None, '', '   ']:
                with self.subTest(field=FIELDS[index], value=value):
                    row = self.valid_row()
                    row[index] = value
                    with self.assertRaisesRegex(ValueError, 'Linha 2:'):
                        load(self.workbook([row]))

    def test_negative_nonfinite_and_invalid_numbers(self):
        for index in [3, 4]:
            for value in [-1, '-0,1', 'NaN', 'Infinity', '-Infinity', 'texto', True, '=1+1']:
                with self.subTest(field=FIELDS[index], value=value):
                    row = self.valid_row()
                    row[index] = value
                    with self.assertRaisesRegex(ValueError, 'Linha 2:'):
                        load(self.workbook([row]))

    def test_category_and_batch_rules(self):
        cases = [['2026-09-21', 'Outra', 'X', 1, 0],
                 ['2026-09-21', 'Ração', 'X', 1.5, 2],
                 ['2026-09-21', 'Recebimentos', 'X', 1, 2]]
        cases += [['2026-09-21', c, 'X', 1, 2] for c in ['Medicamentos', 'Premix', 'Mineral']]
        for row in cases:
            with self.subTest(row=row):
                with self.assertRaisesRegex(ValueError, 'Linha 2:'):
                    load(self.workbook([row]))

    def test_formulas_and_errors_rejected_in_every_column(self):
        for index in range(5):
            for value in ['=1+1', '#DIV/0!']:
                with self.subTest(field=FIELDS[index], value=value):
                    row = self.valid_row()
                    row[index] = value
                    with self.assertRaisesRegex(ValueError, 'Linha 2:'):
                        load(self.workbook([row]))

    def test_blank_rows_skipped_and_physical_line_retained(self):
        rows = [self.valid_row(), [None] * 5, self.valid_row()]
        self.assertEqual(len(load(self.workbook(rows))), 2)
        rows[-1][4] = -1
        with self.assertRaisesRegex(ValueError, r'Linha 4:.*kg'):
            load(self.workbook(rows))

    def test_extra_nonempty_column_rejected(self):
        with self.assertRaisesRegex(ValueError, 'Linha 2:'):
            load(self.workbook([self.valid_row() + ['extra']]))

    def test_first_sheet_used_and_extension_case_insensitive(self):
        path = self.workbook([self.valid_row()], filename='input.XLSX')
        from openpyxl import load_workbook
        wb = load_workbook(path)
        wb.create_sheet('Ignorar').append(['invalido'])
        wb.active = 1
        wb.save(path)
        wb.close()
        self.assertEqual(len(load(path)), 1)

    def test_header_only_workbook_is_empty(self):
        self.assertEqual(load(self.workbook([])), [])

    def test_cli_matches_csv_with_period_and_exports(self):
        raw = [self.valid_row(), ['2026-09-22', 'Ração', 'Produto', 2, '0.2']]
        for suffix, path in [('csv', self.csv_file(raw)), ('xlsx', self.workbook(raw))]:
            result = subprocess.run([sys.executable, str(ROOT / 'src/report.py'), str(path),
                                     '--start', '2026-09-22', '--end', '2026-09-22',
                                     '--output', str(self.folder / suffix)],
                                    capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
        for filename in ['relatorio.html', 'totais.json']:
            self.assertEqual((self.folder / 'csv' / filename).read_bytes(),
                             (self.folder / 'xlsx' / filename).read_bytes())

    def test_cli_corrupt_workbook_has_clean_error_and_no_exports(self):
        path = self.folder / 'corrupt.xlsx'
        path.write_bytes(b'not a workbook')
        output = self.folder / 'result'
        result = subprocess.run([sys.executable, str(ROOT / 'src/report.py'), str(path),
                                 '--output', str(output)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 2)
        self.assertIn('Erro:', result.stderr)
        self.assertNotIn('Traceback', result.stderr)
        self.assertFalse((output / 'totais.json').exists())
        self.assertFalse((output / 'relatorio.html').exists())


if __name__ == '__main__':
    unittest.main()
